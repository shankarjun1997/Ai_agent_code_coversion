"""STM exporter — xlsx + zip output.

Mirrors the artifact's `buildStmWorkbook` + `buildZip`:
  - xlsx with two sheets: STM (one row per mapping) + metadata
  - zip with one .sql per target table + combined .sql + README.txt
"""
from __future__ import annotations

import io
import struct
import zlib
from datetime import datetime, timezone
from typing import List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.stm.blackboard import SqlBundle, StmBlackboard


def build_stm_xlsx(bb: StmBlackboard) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "STM"

    headers = ["Source Table", "Source Column", "Target Table", "Target Column",
               "Mapping Type", "Business Logic"]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFF8E7")
    header_fill = PatternFill(start_color="FF1F2937", end_color="FF1F2937", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for m in bb.mappings:
        ws.append([
            m.source_table, m.source_column,
            m.target_table or "", m.target_column or "",
            m.mapping_type, m.business_logic or "",
        ])

    widths = [26, 30, 28, 28, 14, 80]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    meta = wb.create_sheet("metadata")
    counts = {"1:1": 0, "1:many": 0, "derived": 0, "constant": 0, "unused": 0}
    for m in bb.mappings:
        counts[m.mapping_type] = counts.get(m.mapping_type, 0) + 1
    meta.append(["Generated", datetime.now(timezone.utc).isoformat()])
    meta.append(["Session ID", bb.session_id])
    meta.append(["Source table", bb.source_table_name])
    meta.append(["Target", f"{bb.target_project}.{bb.target_dataset}"])
    meta.append(["Total mappings", len(bb.mappings)])
    for k, v in counts.items():
        meta.append([k, v])
    meta.column_dimensions["A"].width = 18
    meta.column_dimensions["B"].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _crc32(data: bytes) -> int:
    return zlib.crc32(data) & 0xFFFFFFFF


def build_sql_zip(bundle: SqlBundle, source_table_name: str) -> bytes:
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    files: List[Tuple[str, bytes]] = []
    seen: dict[str, int] = {}
    for i, stmt in enumerate(bundle.statements):
        base = stmt.target_table or f"statement_{i + 1:02d}"
        seen[base] = seen.get(base, 0) + 1
        c = seen[base]
        filename = f"{base}.sql" if c == 1 else f"{base}_{c}.sql"
        header = (
            "-- ============================================================\n"
            f"-- Target: {base}\n"
            f"-- Source: {source_table_name}\n"
            f"-- Generated: {today}\n"
            "-- BigQuery Standard SQL\n"
            "-- ============================================================\n\n"
        )
        files.append((filename, (header + stmt.sql + "\n").encode("utf-8")))

    if bundle.combined_sql:
        files.append((f"_all_{source_table_name}_{today}.sql", bundle.combined_sql.encode("utf-8")))

    manifest_lines = [
        f"Source-to-Target Mapping — generated {today}",
        f"Source: {source_table_name}",
        f"Project.Dataset: {bundle.project}.{bundle.dataset}",
        f"Statements: {len(bundle.statements)}",
        "",
        "Files:",
    ]
    for fname, _ in files[:-1]:
        manifest_lines.append(f"  - {fname}")
    if files:
        manifest_lines.append(f"  - {files[-1][0]}   (combined)")
    files.append(("README.txt", "\n".join(manifest_lines).encode("utf-8")))

    return _build_zip(files)


def _build_zip(files: List[Tuple[str, bytes]]) -> bytes:
    local_chunks: List[bytes] = []
    central_chunks: List[bytes] = []
    offset = 0
    for name, data in files:
        name_bytes = name.encode("utf-8")
        crc = _crc32(data)
        size = len(data)
        local_chunks.append(struct.pack(
            "<IHHHHHIIIHH",
            0x04034b50,
            20, 0, 0, 0, 0,
            crc, size, size,
            len(name_bytes), 0,
        ))
        local_chunks.append(name_bytes)
        local_chunks.append(data)
        central_chunks.append(struct.pack(
            "<IHHHHHHIIIHHHHHII",
            0x02014b50,
            20, 20, 0, 0, 0, 0,
            crc, size, size,
            len(name_bytes), 0, 0, 0, 0, 0,
            offset,
        ))
        central_chunks.append(name_bytes)
        offset += 30 + len(name_bytes) + size

    central = b"".join(central_chunks)
    eocd = struct.pack(
        "<IHHHHIIH",
        0x06054b50,
        0, 0, len(files), len(files),
        len(central), offset, 0,
    )
    return b"".join(local_chunks) + central + eocd
