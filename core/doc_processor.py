"""Extract plain text from DOCX, PDF, TXT, or MD requirements documents."""
from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text(path: str) -> str:
    """Return extracted plain text from a requirements document."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _from_pdf(p)
    if suffix in (".docx", ".doc"):
        return _from_docx(p)
    if suffix in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix}. Accepted: pdf, docx, txt, md")


def _from_pdf(path: Path) -> str:
    try:
        import PyPDF2
        parts = []
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    parts.append(text)
        return "\n\n".join(parts)
    except ImportError:
        raise ImportError("PDF support requires: pip install PyPDF2")


def _from_docx(path: Path) -> str:
    try:
        import docx
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = []
        for table in doc.tables:
            for row in table.rows:
                tables.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs + tables)
    except ImportError:
        raise ImportError("DOCX support requires: pip install python-docx")
