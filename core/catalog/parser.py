"""Format-agnostic schema file parser.

Ported from the reference artifact's JS parser
(~/Downloads/Mapping Agent Mock Up/mapping-agent.jsx, lines 74-310).
Three detection strategies, run in priority order:

  Strategy 0  Data sample: header row + many data rows, no table_name column.
              Headers ARE the columns. Types inferred from first ~20 rows.
  Strategy 1  Schema dump: explicit table_name + column_name columns (Task 5).
  Strategy 2  Workbook: one sheet per table (Task 6).

If all strategies fail, raises ParserError. No AI fallback in v1.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl


# ── Type inference constants ────────────────────────────────────────────────

_INT_RE = re.compile(r"^-?\d+$")
_FLOAT_RE = re.compile(r"^-?\d+\.\d+$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}")


# ── Public data structures ──────────────────────────────────────────────────

ParsedColumn = Dict[str, Any]  # {"name": str, "type": str, "description": Optional[str], "ordinal": int}


@dataclass
class ParsedTable:
    table_name: str
    schema_name: Optional[str] = None
    catalog: Optional[str] = None
    description: Optional[str] = None
    columns: List[ParsedColumn] = field(default_factory=list)

    @property
    def column_count(self) -> int:
        return len(self.columns)


@dataclass
class ParsedSchema:
    tables: List[ParsedTable]
    strategy: str  # which strategy succeeded


class ParserError(ValueError):
    """Raised when no strategy can parse the file."""


# ── Header normalization (matches JS norm() at line 80) ─────────────────────

def _norm(s: Any) -> str:
    return re.sub(r"[\s_\-]+", "", str(s or "").lower().strip())


TABLE_KEYS = ["table", "tablename", "table_name", "tbl", "entity", "entity_name",
              "source_table", "target_table", "object_name"]
COLUMN_KEYS = ["column", "columnname", "column_name", "col", "col_name", "field",
               "field_name", "attribute", "columns", "fields"]
TYPE_KEYS = ["type", "datatype", "data_type", "column_type", "field_type",
             "sql_type", "bq_type"]
DESC_KEYS = ["description", "desc", "comment", "comments", "notes", "definition",
             "remarks", "doc"]


def _match_key(headers: List[Any], candidates: List[str]) -> int:
    """Return index of first matching header, or -1 if none. Exact first, then substring."""
    n = [_norm(h) for h in headers]
    c = [_norm(x) for x in candidates]
    for x in c:
        if x in n:
            return n.index(x)
    for x in c:
        for i, h in enumerate(n):
            if x in h:
                return i
    return -1


# ── Type inference (matches JS inferType() at ~line 200) ────────────────────

def _infer_column_type(values: List[Any]) -> str:
    """Inspect up to 20 values, return a single canonical type string."""
    sample = [v for v in values[:20] if v not in (None, "", "NaN")]
    if not sample:
        return "string"

    all_int = True
    all_float = True
    all_date = True
    all_datetime = True
    all_bool = True

    for v in sample:
        s = str(v).strip()
        if isinstance(v, bool):
            all_int = all_float = all_date = all_datetime = False
            continue
        if isinstance(v, int):
            all_float = False
            all_date = all_datetime = all_bool = False
            continue
        if isinstance(v, float):
            all_int = False
            all_date = all_datetime = all_bool = False
            continue
        if isinstance(v, datetime):
            all_int = all_float = all_date = all_bool = False
            continue
        if isinstance(v, date):
            all_int = all_float = all_datetime = all_bool = False
            continue
        # string fallback — try regexes
        if not _INT_RE.match(s):
            all_int = False
        if not _FLOAT_RE.match(s):
            all_float = False
        if not _DATETIME_RE.match(s):
            all_datetime = False
        if not _DATE_RE.match(s):
            all_date = False
        if s.lower() not in ("true", "false", "0", "1"):
            all_bool = False

    if all_datetime:
        return "timestamp"
    if all_date:
        return "date"
    if all_int:
        return "integer"
    if all_float:
        return "float"
    if all_bool:
        return "boolean"
    return "string"


# ── File reading ─────────────────────────────────────────────────────────────

def _read_xlsx(path: Path) -> List[Dict[str, List[List[Any]]]]:
    """Return list of {name, rows2D}, one per sheet."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        out.append({"name": name, "rows2D": rows})
    wb.close()
    return out


# ── Strategy 0: data sample ─────────────────────────────────────────────────

def _try_strategy_0(sheets: List[Dict[str, Any]]) -> Optional[List[ParsedTable]]:
    """Headers ARE column names; rows below are data records."""
    tables: List[ParsedTable] = []
    for sheet in sheets:
        rows = sheet["rows2D"]
        if len(rows) < 6:  # header + 5 data rows minimum
            continue
        headers = rows[0]
        if len(headers) < 2:
            continue

        # Reject if a table_name-style header exists — that's a schema file (strategy 1).
        if _match_key(list(headers), TABLE_KEYS) != -1:
            return None

        # Check distinctness of first column (mostly distinct → data sample).
        first_col_vals = [r[0] for r in rows[1:] if len(r) > 0]
        if not first_col_vals:
            continue
        non_null = [v for v in first_col_vals if v not in (None, "")]
        if not non_null:
            continue
        distinct_ratio = len(set(non_null)) / max(len(non_null), 1)
        if distinct_ratio < 0.7:
            return None

        ident_re = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
        ident_headers = [str(h).strip() for h in headers if h and ident_re.match(str(h).strip())]
        if len(ident_headers) < 2:
            continue

        cols: List[ParsedColumn] = []
        for idx, h in enumerate(headers):
            if not h:
                continue
            col_vals = [r[idx] if idx < len(r) else None for r in rows[1:]]
            cols.append({
                "name": str(h).strip(),
                "type": _infer_column_type(col_vals),
                "description": None,
                "ordinal": idx,
            })

        name = sheet["name"]
        if name.lower() in ("sheet1", "sheet", "(csv)"):
            name = "unknown_table"
        tables.append(ParsedTable(table_name=name, columns=cols))

    return tables or None


# ── Column-list splitting (matches JS splitColumnList at ~line 90) ──────────

def _split_column_list(raw: Any) -> List[str]:
    """Split a comma/newline/pipe/semicolon-delimited column list into parts."""
    s = str(raw or "").strip()
    if not s:
        return []
    # JSON array?
    if s.startswith("["):
        try:
            import json
            arr = json.loads(s)
            if isinstance(arr, list):
                return [str(x).strip() for x in arr if str(x).strip()]
        except ValueError:
            pass
    # Try newline first, then comma/semicolon/pipe. Best split = most parts.
    best = [s]
    for sep_re in (re.compile(r"\r?\n"), re.compile(r"\s*[,;|]\s*")):
        parts = [p.strip() for p in sep_re.split(s) if p.strip()]
        if len(parts) > len(best):
            best = parts
    return best


# ── Strategy 1: explicit schema dump ────────────────────────────────────────

def _try_strategy_1(sheets: List[Dict[str, Any]]) -> Optional[List[ParsedTable]]:
    """Schema-dump file with table_name + column information.

    1a — one row per column: header has table_name AND column_name.
    1b — one row per table: header has table_name AND a list column (`columns`,
         `column_name`, `fields`...). The list column carries CSV-aggregated names.
    """
    for sheet in sheets:
        rows = sheet["rows2D"]
        if not rows:
            continue
        headers = list(rows[0])
        ti = _match_key(headers, TABLE_KEYS)
        ci = _match_key(headers, COLUMN_KEYS)
        if ti == -1 or ci == -1:
            continue
        ty = _match_key(headers, TYPE_KEYS)
        ds = _match_key(headers, DESC_KEYS)

        body = rows[1:]
        if not body:
            continue

        per_table: Dict[str, List[Dict[str, Any]]] = {}
        for r in body:
            if ti >= len(r):
                continue
            tname = r[ti]
            if not tname:
                continue
            tname = str(tname).strip()
            col_cell = r[ci] if ci < len(r) else ""
            type_cell = r[ty] if ty != -1 and ty < len(r) else None
            desc_cell = r[ds] if ds != -1 and ds < len(r) else None

            parts = _split_column_list(col_cell)
            if len(parts) > 1:
                # 1b — explode the CSV into one entry per column.
                for ord_i, p in enumerate(parts):
                    per_table.setdefault(tname, []).append({
                        "name": p, "type": None, "description": None, "ordinal": ord_i,
                    })
            else:
                per_table.setdefault(tname, []).append({
                    "name": str(col_cell).strip(),
                    "type": str(type_cell).strip() if type_cell else None,
                    "description": str(desc_cell).strip() if desc_cell else None,
                    "ordinal": len(per_table.get(tname, [])),
                })

        if not per_table:
            continue
        return [ParsedTable(table_name=tname, columns=cols) for tname, cols in per_table.items()]
    return None


# ── Strategy 2: workbook, one sheet per table ───────────────────────────────

def _try_strategy_2(sheets: List[Dict[str, Any]]) -> Optional[List[ParsedTable]]:
    """Each sheet is a table; the sheet's first row is the column-level header.

    Inside a sheet, header must include a column_name-style header
    (no table_name header — that would belong to strategy 1).
    """
    tables: List[ParsedTable] = []
    for sheet in sheets:
        rows = sheet["rows2D"]
        if len(rows) < 2:
            continue
        headers = list(rows[0])
        # If table_name header is present, this is strategy 1 territory; skip.
        if _match_key(headers, TABLE_KEYS) != -1:
            continue
        ci = _match_key(headers, COLUMN_KEYS)
        if ci == -1:
            continue
        ty = _match_key(headers, TYPE_KEYS)
        ds = _match_key(headers, DESC_KEYS)

        cols: List[ParsedColumn] = []
        for ord_i, r in enumerate(rows[1:]):
            if ci >= len(r):
                continue
            name = r[ci]
            if not name:
                continue
            cols.append({
                "name": str(name).strip(),
                "type": str(r[ty]).strip() if ty != -1 and ty < len(r) and r[ty] else None,
                "description": str(r[ds]).strip() if ds != -1 and ds < len(r) and r[ds] else None,
                "ordinal": ord_i,
            })
        if cols:
            tables.append(ParsedTable(table_name=sheet["name"], columns=cols))
    return tables or None


# ── Main entrypoint ─────────────────────────────────────────────────────────

def parse_schema_file(path: Path) -> ParsedSchema:
    """Parse an xlsx schema or sample file. Raises ParserError on failure."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xls", ".xlsm"):
        sheets = _read_xlsx(path)
    else:
        raise ParserError(f"Unsupported extension {suffix} — only .xlsx for v1")

    s0 = _try_strategy_0(sheets)
    if s0:
        for t in s0:
            if t.table_name == "unknown_table":
                t.table_name = path.stem.replace("_source", "").replace("_sample", "")
        return ParsedSchema(tables=s0, strategy="data_sample")

    s1 = _try_strategy_1(sheets)
    if s1:
        return ParsedSchema(tables=s1, strategy="schema_dump")

    s2 = _try_strategy_2(sheets)
    if s2:
        return ParsedSchema(tables=s2, strategy="workbook_per_table")

    raise ParserError(
        f"No strategy matched {path.name}. Tried: data_sample, schema_dump, workbook_per_table."
    )
