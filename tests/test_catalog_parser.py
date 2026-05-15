"""Parser tests — ported from mapping-agent.jsx tryFastParse strategies."""
from pathlib import Path

import pytest

from core.catalog.parser import ParsedTable, parse_schema_file


FIXTURES = Path(__file__).parent / "fixtures" / "catalog"


def test_strategy_0_data_sample_uses_headers_as_columns():
    """device_activation_event_source.xlsx: 1 sheet with 44 column headers + data rows.
    Strategy 0 should detect data-sample mode and treat headers as the column list.
    """
    result = parse_schema_file(FIXTURES / "device_activation_event_source.xlsx")
    assert len(result.tables) == 1
    t = result.tables[0]
    assert t.table_name == "device_activation_event"
    assert t.column_count == 44
    assert len(t.columns) == 44
    assert any(c["name"] for c in t.columns)
    types = {c.get("type") for c in t.columns}
    assert len(types) > 1, f"All columns inferred as same type — type inference failing: {types}"


def test_strategy_1b_string_agg_schema_dump():
    """Target Information Schema.xlsx: one row per table, with column_name CSV-aggregated.
    Header has table_name + columns (or column_name) + column_count.
    """
    result = parse_schema_file(FIXTURES / "target_information_schema.xlsx")
    # Per the artifact screenshot, this file has 11 tables, 255 columns total.
    assert len(result.tables) == 11
    total_cols = sum(t.column_count for t in result.tables)
    assert total_cols == 255

    # Spot-check known tables from the screenshot.
    names = {t.table_name for t in result.tables}
    assert "billing_invoice" in names
    assert "customer_account_profile" in names
    assert "network_outage_event" in names

    # Verify column counts on a few from the screenshot.
    by_name = {t.table_name: t for t in result.tables}
    assert by_name["billing_invoice"].column_count == 24
    assert by_name["customer_account_profile"].column_count == 24
    assert by_name["network_outage_event"].column_count == 23


def test_strategy_2_workbook_per_table(tmp_path):
    """Synthetic workbook: 2 sheets, each a table with column rows."""
    import openpyxl
    wb = openpyxl.Workbook()
    s1 = wb.active
    s1.title = "orders"
    s1.append(["column_name", "data_type", "description"])
    s1.append(["order_id", "STRING", "PK"])
    s1.append(["amount", "FLOAT64", "USD"])
    s2 = wb.create_sheet("customers")
    s2.append(["column_name", "data_type", "description"])
    s2.append(["customer_id", "STRING", "PK"])
    s2.append(["email", "STRING", None])
    s2.append(["created_at", "TIMESTAMP", None])

    path = tmp_path / "workbook.xlsx"
    wb.save(path)

    result = parse_schema_file(path)
    assert result.strategy == "workbook_per_table"
    by_name = {t.table_name: t for t in result.tables}
    assert set(by_name.keys()) == {"orders", "customers"}
    assert by_name["orders"].column_count == 2
    assert by_name["customers"].column_count == 3
    assert by_name["orders"].columns[0]["name"] == "order_id"
    assert by_name["orders"].columns[0]["type"] == "STRING"
